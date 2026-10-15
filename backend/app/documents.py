from __future__ import annotations
import io
import json
from fastapi import APIRouter,Depends,HTTPException
from fastapi.responses import StreamingResponse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from backend.app import db
from backend.app.security import current_user,require_operator
router=APIRouter(prefix='/api/documents')

def get_document(did):
    doc=db.one('SELECT * FROM documents WHERE id=?',(did,))
    if not doc:raise HTTPException(404,'文书不存在')
    doc['content']=json.loads(doc['content_json']);doc['reviews']=db.rows('SELECT * FROM document_reviews WHERE document_id=? ORDER BY id',(did,))
    return doc

@router.get('/{did:int}')
def detail(did:int,user=Depends(current_user)):return {'code':0,'data':get_document(did),'message':'ok'}

@router.post('/{did:int}/save-draft')
def save(did:int,user=Depends(current_user)):
    require_operator(user);d=get_document(did)
    if d['status']=='issued':raise ValueError('已签发文书不可覆盖为草稿')
    with db.connect() as c:db.audit(c,user['id'],'save_draft','document',did)
    return {'code':0,'data':{'id':did,'status':d['status']},'message':'已保存，待人工审核'}

def word_bytes(d):
    doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11);section.top_margin=section.bottom_margin=Inches(.8)
    style=doc.styles['Normal'];style.font.name='宋体';style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');style.font.size=Pt(11);style.paragraph_format.space_after=Pt(7)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0);doc.styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
    doc.add_paragraph(d['title'],style='Title')
    doc.add_paragraph('状态：'+('已人工签发' if d['status']=='issued' else '草稿 待人工审核签字后生效'))
    if d.get('doc_no'):doc.add_paragraph('文书编号：'+d['doc_no'])
    content=d['content']
    for key,label in [('hazard_title','隐患事项'),('location','现场位置'),('responsible_team','责任班组'),('responsible_person','责任人'),('deadline','整改期限')]:
        if key in content:doc.add_paragraph(label+'：'+str(content[key] or '待人工补充'))
    body=content.get('body','')
    if body:
        doc.add_paragraph('办理内容',style='Heading 1')
        for line in str(body).splitlines():
            if line.strip():doc.add_paragraph(line)
    evidence=content.get('evidence',[])
    if evidence:
        doc.add_paragraph('原文依据',style='Heading 1')
        for e in evidence:
            doc.add_paragraph(f"{e.get('name','')} {e.get('doc_no','')} {e.get('section','')} 第{e.get('page',1)}页")
            doc.add_paragraph(e.get('text',''))
            doc.add_paragraph('片段：'+e.get('chunk_id',''))
    if d['status']=='issued':
        u=db.one('SELECT display_name FROM users WHERE id=?',(d['issued_by'],))
        doc.add_paragraph('人工签发：'+(u['display_name'] if u else str(d['issued_by']))+'  '+d['issued_at'])
    else:doc.add_paragraph('人工审核人及日期：________________')
    for tree in (doc._element,doc.styles.element):
        for border in list(tree.xpath('.//w:pBdr')):border.getparent().remove(border)
    buffer=io.BytesIO();doc.save(buffer);return buffer.getvalue()

@router.get('/{did:int}/export')
def export(did:int,user=Depends(current_user)):
    d=get_document(did);data=word_bytes(d)
    with db.connect() as c:db.audit(c,user['id'],'export_word','document',did)
    return StreamingResponse(io.BytesIO(data),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="document-{did}.docx"'})
