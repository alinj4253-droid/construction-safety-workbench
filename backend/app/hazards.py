from __future__ import annotations
import uuid
import json
from datetime import timedelta,datetime
from pathlib import Path
from fastapi import APIRouter,Depends,HTTPException,UploadFile,File,Form
from fastapi.responses import StreamingResponse,FileResponse
from pydantic import BaseModel,Field
from backend.app import db
from backend.app.security import current_user,require_operator

router=APIRouter(prefix='/api/hazards')
def decorated(h):return {**h,'is_demo':bool(h.get('is_demo',False) or (h.get('code') or '').startswith('DEMO-')),'overdue':db.overdue(h),'overdue_days':max(0,(db.today()-__import__('datetime').date.fromisoformat(h['deadline'][:10])).days) if h.get('deadline') and db.overdue(h) else 0}

HAZARD_SELECT='SELECT h.*,CASE WHEN d.hazard_id IS NULL THEN 0 ELSE 1 END is_demo FROM hazards h LEFT JOIN hazard_demo_cases d ON d.hazard_id=h.id'

def get_hazard(hid):
    h=db.one(HAZARD_SELECT+' WHERE h.id=?',(hid,))
    if not h:raise HTTPException(404,'隐患不存在')
    return decorated(h)

@router.get('')
def list_hazards(filter:str='all',user=Depends(current_user)):
    all_rows=[decorated(h) for h in db.rows(HAZARD_SELECT+' ORDER BY h.id DESC')]
    selected=[h for h in all_rows if filter=='all' or h['track']==filter or filter in {'over','overdue'} and h['overdue']]
    return {'code':0,'data':{'items':selected,'counts':{'all':len(all_rows),'light':sum(h['track']=='light' for h in all_rows),'heavy':sum(h['track']=='heavy' for h in all_rows),'over':sum(h['overdue'] for h in all_rows)}},'message':'ok'}

class HazardInput(BaseModel):
    title:str=Field(min_length=2,max_length=300)
    category:str='现场隐患'
    location:str=''
    description:str=''
    track:str='heavy'
    source:str='manual'
    responsible_team:str=''
    responsible_person:str=''
    deadline:str|None=None

@router.post('')
def create_hazard(body:HazardInput,user=Depends(current_user)):
    require_operator(user)
    if body.track not in {'light','heavy'}:raise ValueError('请选择轻轨或重轨')
    if body.source!='manual':raise ValueError('监控隐患必须从抓拍记录人工确认，不能绕过确认入口')
    if body.deadline:
        __import__('datetime').date.fromisoformat(body.deadline)
    with db.connect() as c:
        cur=c.execute('INSERT INTO hazards(code,title,category,location,description,track,level,status,source,responsible_team,responsible_person,deadline,created_by,created_at,confirmed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
          ('HZ-'+db.today().strftime('%Y%m%d')+'-'+uuid.uuid4().hex[:8].upper(),body.title,body.category,body.location,body.description or body.title,body.track,'一般隐患' if body.track=='heavy' else '轻微违章','待整改' if body.track=='heavy' else '待确认','manual',body.responsible_team,body.responsible_person,body.deadline,user['id'],db.now(),db.now() if body.track=='heavy' else None))
        hid=cur.lastrowid;db.event(c,hid,'register',user['id'],body.model_dump())
    return {'code':0,'data':get_hazard(hid),'message':'ok'}

class DemoInput(BaseModel):confirmed:bool=False

@router.post('/demo-cases')
def prepare_demo(body:DemoInput,user=Depends(current_user)):
    require_operator(user)
    if not body.confirmed:raise ValueError('请确认创建新的演示模拟案例，原记录不会被重置')
    title='3#楼东侧基坑临边防护缺失，未设防护栏杆'
    reply='【演示模拟】木工班组已完成3#楼东侧基坑临边防护整改，补设防护栏杆及挡脚板，现申请安全员复查。'
    with db.connect() as c:
        stamp=db.now()
        hid=c.execute('INSERT INTO hazards(code,title,category,location,description,track,level,status,source,responsible_team,responsible_person,deadline,created_by,created_at,confirmed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',('DEMO-'+db.today().strftime('%Y%m%d')+'-'+uuid.uuid4().hex[:8].upper(),title,'临边防护','3#楼东侧基坑','【演示模拟】用于答辩闭环演示，不代表实际现场事实。'+title,'heavy','一般隐患','待整改','manual','木工班组','木工班组负责人（演示角色）',str(db.today()+timedelta(days=3)),user['id'],stamp,stamp)).lastrowid
        c.execute('INSERT INTO hazard_demo_cases VALUES(?,?,?,?,?)',(hid,reply,'木工班组',user['id'],stamp))
        db.event(c,hid,'prepare_demo_case',user['id'],{'is_simulated':True,'note':'新建演示记录，未重置或删除任何原记录；预置模拟反馈尚未接收登记'})
    return {'code':0,'data':get_hazard(hid),'message':'已新建演示模拟案例，原档案保留'}

@router.get('/{hid:int}')
def detail(hid:int,user=Depends(current_user)):
    h=get_hazard(hid)
    h['events']=db.rows('SELECT e.*,u.display_name actor_name FROM hazard_events e LEFT JOIN users u ON CAST(u.id AS TEXT)=e.actor WHERE e.hazard_id=? ORDER BY e.id',(hid,))
    h['documents']=db.rows('SELECT id,doc_no,title,status,issued_at FROM documents WHERE hazard_id=? ORDER BY id DESC',(hid,))
    h['attachments']=db.rows('SELECT a.*,u.display_name recorded_by_name FROM attachments a LEFT JOIN users u ON u.id=a.created_by WHERE a.hazard_id=? ORDER BY a.id',(hid,))
    for a in h['attachments']:
        a['is_image']=Path(a['file_path']).suffix.lower() in {'.png','.jpg','.jpeg'}
        a['available']=(db.ROOT/a.pop('file_path')).is_file()
    h['feedbacks']=db.rows('SELECT f.*,u.display_name recorded_by_name FROM hazard_feedbacks f LEFT JOIN users u ON u.id=f.recorded_by WHERE f.hazard_id=? ORDER BY f.id DESC',(hid,))
    h['demo_feedback']=db.one('SELECT feedback_content,feedback_team,prepared_at FROM hazard_demo_cases WHERE hazard_id=?',(hid,))
    if h['demo_feedback']:h['demo_feedback']['received']=any(f['feedback_source']=='demo_simulation' for f in h['feedbacks'])
    if not h['feedbacks'] and h['rectify_note']:
        event=next((e for e in reversed(h['events']) if e['action']=='submit_rectify' and json.loads(e['detail']).get('note')==h['rectify_note']),None)
        h['feedbacks']=[{'id':None,'content':h['rectify_note'],'responsible_team':h['responsible_team'],'feedback_source':'legacy','provided_by':None,'provided_at':None,'recorded_by_name':event.get('actor_name') if event else None,'recorded_at':event['created_at'] if event else None,'is_simulated':h['is_demo']}]
    return {'code':0,'data':h,'message':'ok'}

class ActionInput(BaseModel):
    action:str
    note:str=''
    version:int|None=None
    feedback_source:str|None=None
    feedback_team:str|None=Field(default=None,max_length=120)
    feedback_provider:str=Field(default='',max_length=120)
    feedback_at:datetime|None=None

@router.post('/{hid:int}/actions')
def act(hid:int,body:ActionInput,user=Depends(current_user)):
    if body.action!='urge':require_operator(user)
    with db.connect() as c:
        c.execute('BEGIN IMMEDIATE')
        row=c.execute('SELECT * FROM hazards WHERE id=?',(hid,)).fetchone()
        if not row:raise HTTPException(404,'隐患不存在')
        h=dict(row)
        if body.version is not None and body.version!=h['version']:raise HTTPException(409,'记录已更新，请刷新后操作')
        note=body.note
        feedback=None
        if body.action=='submit_rectify' and h['track']=='heavy':
            source=body.feedback_source or 'safety_officer_recorded'
            if source not in {'safety_officer_recorded','demo_simulation'}:raise ValueError('反馈必须由安全员登记，不能伪造班组账号提交')
            demo=c.execute('SELECT * FROM hazard_demo_cases WHERE hazard_id=?',(hid,)).fetchone()
            team=(body.feedback_team if body.feedback_team is not None else h['responsible_team']) or ''
            if body.feedback_team is not None and not team.strip():raise ValueError('请填写反馈所属责任班组')
            if h['responsible_team'] and team.strip()!=h['responsible_team'].strip():raise ValueError('反馈班组必须与当前隐患责任班组一致')
            if source=='demo_simulation':
                if not demo:raise ValueError('普通隐患没有预置演示反馈')
                if c.execute("SELECT id FROM hazard_feedbacks WHERE hazard_id=? AND feedback_source='demo_simulation'",(hid,)).fetchone():raise ValueError('该模拟反馈已登记，退回复查后请代录新的反馈或新建演示案例')
                if demo['feedback_team']!=h['responsible_team']:raise ValueError('责任班组已变化，不能使用原模拟反馈')
                note=demo['feedback_content'];team=demo['feedback_team']
            feedback={'feedback_source':source,'responsible_team':team.strip(),'provided_by':body.feedback_provider.strip() if source!='demo_simulation' else None,'provided_at':body.feedback_at.isoformat() if body.feedback_at and source!='demo_simulation' else None,'is_simulated':bool(demo) or h['code'].startswith('DEMO-'),'recording_mode':'安全员登记演示模拟反馈' if source=='demo_simulation' else '安全员代录班组反馈'}
        status=db.next_status(h['track'],h['status'],body.action,note)
        if feedback:
            feedback['feedback_id']=c.execute('INSERT INTO hazard_feedbacks(hazard_id,content,responsible_team,feedback_source,provided_by,provided_at,recorded_by,recorded_at,is_simulated) VALUES(?,?,?,?,?,?,?,?,?)',(hid,note,feedback['responsible_team'],feedback['feedback_source'],feedback['provided_by'],feedback['provided_at'],user['id'],db.now(),int(feedback['is_simulated']))).lastrowid
        c.execute('UPDATE hazards SET status=?,version=version+1,rectify_note=CASE WHEN ?=\'submit_rectify\' THEN ? ELSE rectify_note END,review_note=CASE WHEN ? IN (\'review_close\',\'review_return\') THEN ? ELSE review_note END,confirmed_at=CASE WHEN ?=\'confirm\' THEN ? ELSE confirmed_at END,closed_at=CASE WHEN ? IN (\'已闭环\',\'已销项\') THEN ? ELSE closed_at END WHERE id=?',
                  (status,body.action,note,body.action,note,body.action,db.now(),status,db.now(),hid))
        db.event(c,hid,body.action,user['id'],{'before':h['status'],'after':status,'note':note,'delivery':'仅记录催办，未对接短信' if body.action=='urge' else None,**(feedback or {})})
    return {'code':0,'data':get_hazard(hid),'message':'ok'}

@router.get('/export')
def export_hazards(filter:str='all',format:str='csv',user=Depends(current_user)):
    import io,csv,json
    records=list_hazards(filter,user)['data']['items']
    fields=list(records[0]) if records else ['id','code','title','track','status','source','deadline']
    fields+=['documents','citations','event_summary','attachments']
    result=[]
    for h in records:
        docs=db.rows('SELECT id,doc_no,status FROM documents WHERE hazard_id=?',(h['id'],))
        citations=db.rows('SELECT dc.chunk_id,dc.snapshot_json FROM document_citations dc JOIN documents d ON d.id=dc.document_id WHERE d.hazard_id=?',(h['id'],))
        events=db.rows('SELECT actor,action,detail,created_at FROM hazard_events WHERE hazard_id=? ORDER BY id',(h['id'],))
        attachments=db.rows('SELECT id,kind,original_name FROM attachments WHERE hazard_id=?',(h['id'],))
        row={**h,'documents':db.dumps(docs),'citations':db.dumps(citations),'event_summary':db.dumps(events),'attachments':db.dumps(attachments)}
        result.append([row.get(key,'') for key in fields])
    def safe(value):
        s=str(value if value is not None else '')
        return "'"+s if s.startswith(('=','+','-','@','\t','\r')) else s
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        book=Workbook();sheet=book.active;sheet.title='隐患迎检台账';sheet.append(fields)
        for row in result:sheet.append([safe(v) for v in row])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17324D')
        buf=io.BytesIO();book.save(buf);payload=buf.getvalue();media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet';ext='xlsx'
    elif format=='csv':
        out=io.StringIO();writer=csv.writer(out);writer.writerow(fields);writer.writerows([[safe(v) for v in row] for row in result]);payload=out.getvalue().encode('utf-8-sig');media='text/csv;charset=utf-8';ext='csv'
    else:raise ValueError('仅支持CSV或Excel')
    with db.connect() as c:db.audit(c,user['id'],'export_hazards','hazard','all',{'filter':filter,'count':len(records),'format':format})
    return StreamingResponse(io.BytesIO(payload),media_type=media,headers={'Content-Disposition':f'attachment; filename="hazards-{db.today()}.{ext}"'})

@router.post('/{hid:int}/attachments')
async def add_attachment(hid:int,file:UploadFile=File(...),kind:str=Form('rectify'),user=Depends(current_user)):
    import hashlib
    from pathlib import Path
    require_operator(user);h=get_hazard(hid)
    if h['status'] in db.TERMINAL:raise ValueError('闭环档案只读，不能再追加整改证据')
    if kind not in {'before','rectify','review'}:raise ValueError('附件类型无效')
    data=await file.read(10*1024*1024+1)
    if len(data)>10*1024*1024:raise ValueError('单个附件最大10MB')
    ext='.png' if data.startswith(b'\x89PNG\r\n\x1a\n') else '.jpg' if data.startswith(b'\xff\xd8\xff') else '.pdf' if data.startswith(b'%PDF') else None
    if not ext:raise ValueError('仅接受有效PNG、JPEG或PDF证据')
    directory=db.ROOT/'backend/data/attachments';directory.mkdir(parents=True,exist_ok=True)
    path=directory/(uuid.uuid4().hex+ext);path.write_bytes(data)
    with db.connect() as c:
        aid=c.execute('INSERT INTO attachments(hazard_id,kind,file_path,original_name,created_by,created_at) VALUES(?,?,?,?,?,?)',(hid,kind,str(path.relative_to(db.ROOT)),Path(file.filename or '附件').name,user['id'],db.now())).lastrowid
        db.event(c,hid,'upload_evidence',user['id'],{'attachment_id':aid,'kind':kind,'sha256':hashlib.sha256(data).hexdigest()})
    return {'code':0,'data':{'id':aid},'message':'证据已保存并留痕'}

@router.get('/attachments/{aid:int}')
def attachment_file(aid:int,user=Depends(current_user)):
    a=db.one('SELECT * FROM attachments WHERE id=?',(aid,))
    if not a or not (db.ROOT/a['file_path']).is_file():raise HTTPException(404,'附件不存在或文件缺失')
    return FileResponse(db.ROOT/a['file_path'],filename=a['original_name'])
